from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"C:\Users\USER\Desktop\Quiks")
MEDIA = ROOT / ".tmp" / "proposal-unzip" / "word" / "media"
OUT = ROOT / "deliverables" / "Quiks_School_Partnership_Proposal_Refined.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)

BLACK = "000000"
TEXT = "273142"
MUTED = "526273"
PURPLE = "6F2DBD"
NAVY = "17324D"
PALE = "F4F1FA"
PALE_BLUE = "F2F7FA"
WHITE = "FFFFFF"
BORDER = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def set_font(run, name="Aptos", size=None, bold=None, color=TEXT, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, before=0, after=6, line=1.12, keep=False):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep


def add_text(doc, text, style=None, bold_lead=None, after=6):
    p = doc.add_paragraph(style=style)
    style_paragraph(p, after=after)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_bullet(doc, lead, body, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    style_paragraph(p, after=4)
    r = p.add_run(lead)
    set_font(r, bold=True)
    r = p.add_run(body)
    set_font(r)
    return p


def add_number(doc, lead, body):
    p = doc.add_paragraph(style="List Number")
    style_paragraph(p, after=5)
    r = p.add_run(lead)
    set_font(r, bold=True)
    r = p.add_run(body)
    set_font(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_font(r, size=18 if level == 1 else 13, bold=True, color=BLACK)
    return p


def add_section_marker(doc, number, label):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{number}  |  {label.upper()}")
    set_font(r, size=9, bold=True, color=PURPLE)
    return p


def add_page_break(doc):
    doc.add_page_break()


def add_table_text(cell, text, bold=False, color=TEXT, size=9.2, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_end])
    set_font(run, size=8, color=MUTED)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.68)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(TEXT)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12

title = styles["Title"]
title.font.name = "Aptos Display"
title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
title.font.size = Pt(30)
title.font.bold = True
title.font.color.rgb = RGBColor.from_string(BLACK)
title_ppr = title._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

for n, size in ((1, 18), (2, 13)):
    s = styles[f"Heading {n}"]
    s.font.name = "Aptos Display"
    s._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    s._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor.from_string(BLACK)
    s.paragraph_format.keep_with_next = True

for list_name in ("List Bullet", "List Bullet 2", "List Number"):
    s = styles[list_name]
    s.font.name = "Aptos"
    s._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    s._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    s.font.size = Pt(10.3)
    s.font.color.rgb = RGBColor.from_string(TEXT)

# Footer
footer = section.footer
fp = footer.paragraphs[0]
fp.text = "Quiks School Partnership Proposal    |    "
for run in fp.runs:
    set_font(run, size=8, color=MUTED)
add_page_number(fp)

# Cover
p = doc.add_paragraph(style="Title")
remove_paragraph_borders(p)
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.paragraph_format.space_before = Pt(34)
p.paragraph_format.space_after = Pt(10)
r = p.add_run("Quiks School Partnership Proposal")
set_font(r, name="Aptos Display", size=30, bold=True, color=BLACK)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(20)
r = p.add_run("A structured digital learning partnership for students, teachers and school leadership")
set_font(r, size=14, color=PURPLE)

meta = doc.add_table(rows=3, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.LEFT
meta.autofit = False
for row in meta.rows:
    keep_row_together(row)
labels = ["Prepared for", "Prepared by", "Date"]
values = ["[SCHOOL NAME]", "Tech Solution Providers Ltd", "[DATE]"]
for i, (label, value) in enumerate(zip(labels, values)):
    set_width(meta.cell(i, 0), 1.35)
    set_width(meta.cell(i, 1), 5.45)
    add_table_text(meta.cell(i, 0), label.upper(), bold=True, color=PURPLE, size=8.5)
    add_table_text(meta.cell(i, 1), value, bold=(i == 0), size=10.5)
for row in meta.rows:
    for cell in row.cells:
        set_cell_shading(cell, WHITE)
set_table_borders(meta, color=WHITE, size="0")

doc.add_paragraph().paragraph_format.space_after = Pt(12)
art = doc.add_table(rows=1, cols=3)
art.alignment = WD_TABLE_ALIGNMENT.CENTER
art.autofit = False
for i, (image_name, label) in enumerate((("image1.png", "QUIKS CHILDREN"), ("image2.png", "QUIKS TEENS"), ("image3.png", "QUIKS UNI"))):
    cell = art.cell(0, i)
    set_width(cell, 2.2)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    picture = p.add_run().add_picture(str(MEDIA / image_name), width=Inches(1.45))
    picture._inline.docPr.set("descr", f"{label.title()} learning platform illustration")
    picture._inline.docPr.set("title", label.title())
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(0)
    r = p2.add_run(label)
    set_font(r, size=8.5, bold=True, color=PURPLE)
set_table_borders(art, color=WHITE, size="0")

doc.add_paragraph().paragraph_format.space_after = Pt(10)
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.space_after = Pt(6)
r = p.add_run("Partnership purpose")
set_font(r, size=12, bold=True, color=BLACK)
add_text(doc, "To give students structured practice and learning support beyond lesson time, while equipping teachers and school leaders with practical tools for classroom delivery, assessment, record management and academic follow-up.", after=0)

# Page 2
add_page_break(doc)
add_section_marker(doc, "1", "Overview")
add_heading(doc, "Why Quiks", 1)
add_text(doc, "Quiks is a learning and practice platform designed for different stages of education. Quiks Children supports younger learners, Quiks Teens serves secondary-school and college-age students, and Quiks Uni is designed for tertiary learners. A school can adopt the variant or combination of variants that fits its age groups, curriculum needs and examination priorities.")
add_text(doc, "The platform complements classroom teaching. Students can practise across subjects, revise selected topics, receive explanations and use the Learning Hub for further study. Teachers can organise classes, publish lesson notes, create tests and assignments, review questions before release, communicate with learners and monitor submitted work. School administrators can manage enrolment, classes, licences and consolidated activity records from the school portal.")
add_text(doc, "Quiks is available on web and mobile. Signed-in users can access the same school membership and learning records across supported devices within the appropriate Quiks variant. School packages are provided without advertising.")

add_heading(doc, "What the school receives", 2)
add_bullet(doc, "A school administration portal. ", "Authorised administrators can manage school enrolment, roles, licence dates, class links and central activity records.")
add_bullet(doc, "Flexible enrolment. ", "Schools may use a shared enrolment code or issue unique codes to individual students and staff. The school can configure the profile fields it requires for enrolment.")
add_bullet(doc, "Connected classrooms. ", "Classes may be created centrally by an administrator or created by teachers and linked to the school portal, keeping classroom activities and results associated with the school.")
add_bullet(doc, "Digital assessment support. ", "Teachers can prepare tests and assignments for classroom use and computer-based testing, set deadlines, review generated questions and receive objective-question scores for review.")
add_bullet(doc, "Teaching and learning tools. ", "Lesson notes, class communication, practice sessions, the Learning Hub and Competition Arena support both directed and independent learning.")
add_bullet(doc, "Cross-device access. ", "Students and staff can use supported mobile or web access without a school having to maintain a separate app for each campus.")

add_heading(doc, "A practical role in school delivery", 2)
add_text(doc, "A school can use Quiks for homework, revision, intervention, enrichment, lesson-note distribution, continuous assessment and selected CBT activities. Teachers retain responsibility for curriculum delivery, question approval and the interpretation of results, particularly where an activity contributes to a high-stakes decision.")

# Page 3
add_page_break(doc)
add_section_marker(doc, "2", "Value to stakeholders")
add_heading(doc, "Value to the school", 1)
add_bullet(doc, "Central academic visibility. ", "The school portal brings together authorised enrolment records, linked classes and submitted activity results so administrators can review participation and performance across the school.")
add_bullet(doc, "A structured CBT pathway. ", "Schools can use teacher-reviewed tests and assignments for routine assessments and CBT practice, reducing dependence on separate systems for basic digital assessment workflows.")
add_bullet(doc, "Consistent learning beyond school hours. ", "Class activities, lesson notes and practice sessions give students an organised route to continue learning at home.")
add_bullet(doc, "Scalable access. ", "Licence bands support small schools, growing schools, larger enrolments and multi-campus school groups.")
add_bullet(doc, "Operational control. ", "Role-based school administration, enrolment codes, licence validity dates and configurable enrolment fields allow each school to control who joins and what information is collected.")
add_bullet(doc, "Clearer performance evidence. ", "Central results and classroom records can support follow-up, intervention planning and communication with parents, subject to the school's assessment and reporting processes.")
add_bullet(doc, "Commercial flexibility. ", "A school may sponsor access, include it in an approved technology or learning fee, or use the per-learner option, as agreed in the final commercial arrangement.")

add_heading(doc, "Value to teachers", 1)
add_bullet(doc, "Faster activity preparation. ", "Teachers can create tests and assignments by subject, topic, difficulty, question count and deadline, then review and edit the questions before publishing.")
add_bullet(doc, "Lesson-note support. ", "Teachers can upload notes as prepared or request minimal, substantial or deep refinement. The refined note remains available for teacher review and editing before publication.")
add_bullet(doc, "Controlled access to notes. ", "A teacher can publish a note as read-only or allow students to download it, and can later edit or delete the note with confirmation.")
add_bullet(doc, "Activities based on lesson content. ", "A teacher can generate a test or assignment from a lesson note, helping assessment stay connected to material already taught.")
add_bullet(doc, "Reduced marking workload. ", "Objective questions can be scored automatically, allowing teachers to spend more time reviewing misconceptions and planning follow-up instruction.")
add_bullet(doc, "Class communication. ", "The shared class chat allows teachers and enrolled students to exchange messages within the classroom space.")
add_bullet(doc, "Accessible teaching resources. ", "The Learning Hub supports research, lesson preparation and classroom enrichment across subjects and levels.")

# Page 4
add_page_break(doc)
add_section_marker(doc, "2", "Value to stakeholders continued")
add_heading(doc, "Value to students", 1)
add_bullet(doc, "Guided independent study. ", "Students can practise broadly across a subject or focus on selected topics, with difficulty and level settings suited to their learning stage.")
add_bullet(doc, "Immediate learning support. ", "Explanations and the Learning Hub help students understand why an answer is correct and study a concept in more depth.")
add_bullet(doc, "Exam-focused preparation. ", "Students can practise within relevant grade, curriculum and target-examination contexts.")
add_bullet(doc, "Classroom continuity. ", "Learners can receive lesson notes, complete tests and assignments, view permitted resources and communicate within their class from supported devices.")
add_bullet(doc, "Motivation and participation. ", "The Competition Arena and Top Performers view provide an additional way to practise, compare performance and represent the learner's school where a school name is recorded.")
add_bullet(doc, "Study habit tracking. ", "Daily study targets and recorded learning time help students build a more consistent practice routine. Time spent on classroom tests and assignments contributes to the learner's study activity.")

add_heading(doc, "Value to parents and guardians", 1)
add_bullet(doc, "A structured home-study option. ", "Quiks gives families an organised learning resource for practice, revision and assigned work outside school hours.")
add_bullet(doc, "Better continuity with schoolwork. ", "Published lesson notes and classroom activities help parents understand what the learner has been asked to study or complete.")
add_bullet(doc, "Visible study habits. ", "The learner's study target, time spent and performance records can support practical conversations about consistency and areas requiring attention.")
add_bullet(doc, "Lower school-channel pricing. ", "School and per-learner packages provide a lower-cost route than purchasing the standard individual subscription separately, depending on the selected package and enrolled population.")
add_text(doc, "Quiks does not replace parental supervision or teacher judgement. Schools and families should agree appropriate expectations for devices, screen time, acceptable use and AI-assisted learning.", after=0)

# Page 5 financials
add_page_break(doc)
add_section_marker(doc, "3", "Financial implications")
add_heading(doc, "Pricing and licence options", 1)
add_text(doc, "The school may begin with a four-week pilot at no licence cost. A pilot extension of up to eight weeks may be agreed where the school calendar or evaluation plan requires more time. If the school decides not to proceed after the agreed pilot, no onboarding or licence fee is payable.")
add_text(doc, "If the school adopts Quiks, a one-time onboarding fee of NGN 50,000 becomes payable. The school then selects the licence option that matches its enrolment and preferred payment cycle.")

table = doc.add_table(rows=1, cols=5)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [1.35, 2.05, 1.10, 1.15, 1.15]
headers = ["Licence option", "Eligible users", "Monthly fee", "Per term", "Per session"]
for i, (head, width) in enumerate(zip(headers, widths)):
    set_width(table.cell(0, i), width)
    set_cell_shading(table.cell(0, i), NAVY)
    add_table_text(table.cell(0, i), head, bold=True, color=WHITE, size=8.7, align=WD_ALIGN_PARAGRAPH.CENTER)
set_repeat_table_header(table.rows[0])

rows = [
    ("Per Learner Access", "Individual students enrolled through the school", "NGN 1,500 per learner", "Approx. NGN 4,500 per learner", "By agreement"),
    ("Essential School", "10 to 100 students", "Not applicable", "NGN 300,000", "NGN 750,000"),
    ("Growth School", "101 to 200 students", "Not applicable", "NGN 500,000", "NGN 1,250,000"),
    ("Comprehensive School", "201 to 500 students", "Not applicable", "NGN 800,000", "NGN 1,950,000"),
    ("Enterprise Network", "A group of schools or multiple campuses; enrolment scope confirmed in the agreement", "Not applicable", "NGN 1,000,000", "NGN 2,500,000"),
]
for ri, data in enumerate(rows, start=1):
    cells = table.add_row().cells
    keep_row_together(table.rows[-1])
    for ci, (value, width) in enumerate(zip(data, widths)):
        set_width(cells[ci], width)
        set_cell_shading(cells[ci], WHITE if ri % 2 else PALE_BLUE)
        align = WD_ALIGN_PARAGRAPH.LEFT if ci in (0, 1) else WD_ALIGN_PARAGRAPH.CENTER
        add_table_text(cells[ci], value, bold=(ci == 0), size=8.4, align=align)
set_table_borders(table)

add_heading(doc, "Session payment savings", 2)
add_text(doc, "A session licence costs less than purchasing three separate term licences:")
savings = doc.add_table(rows=1, cols=4)
savings.alignment = WD_TABLE_ALIGNMENT.CENTER
savings.autofit = False
sw = [2.10, 1.45, 1.45, 1.45]
for i, h in enumerate(("Licence option", "Three terms", "Session fee", "Saving")):
    set_width(savings.cell(0, i), sw[i])
    set_cell_shading(savings.cell(0, i), PURPLE)
    add_table_text(savings.cell(0, i), h, bold=True, color=WHITE, size=8.8, align=WD_ALIGN_PARAGRAPH.CENTER)
set_repeat_table_header(savings.rows[0])
for ri, data in enumerate((
    ("Essential School", "NGN 900,000", "NGN 750,000", "NGN 150,000"),
    ("Growth School", "NGN 1,500,000", "NGN 1,250,000", "NGN 250,000"),
    ("Comprehensive School", "NGN 2,400,000", "NGN 1,950,000", "NGN 450,000"),
    ("Enterprise Network", "NGN 3,000,000", "NGN 2,500,000", "NGN 500,000"),
), start=1):
    cells = savings.add_row().cells
    keep_row_together(savings.rows[-1])
    for ci, (value, width) in enumerate(zip(data, sw)):
        set_width(cells[ci], width)
        set_cell_shading(cells[ci], WHITE if ri % 2 else PALE)
        add_table_text(cells[ci], value, bold=(ci in (0, 3)), size=8.6, align=WD_ALIGN_PARAGRAPH.CENTER if ci else WD_ALIGN_PARAGRAPH.LEFT)
set_table_borders(savings)

# Page 6 commercial detail
add_page_break(doc)
add_section_marker(doc, "3", "Financial implications continued")
add_heading(doc, "What the licence price covers", 1)
add_bullet(doc, "Platform access. ", "Use of the relevant Quiks variant on supported web and mobile devices for the licensed users and validity period.")
add_bullet(doc, "School administration. ", "Access to school enrolment, role management, class linkage, licence information and central activity records.")
add_bullet(doc, "Teaching and assessment features. ", "Classroom management, lesson notes, tests, assignments, class chat, practice sessions, the Learning Hub and Competition Arena, subject to the package and product availability.")
add_bullet(doc, "Initial onboarding support. ", "Account setup guidance, school administrator orientation, enrolment configuration and a short teacher introduction are covered by the one-time onboarding fee.")
add_bullet(doc, "Advertising-free school use. ", "School-licensed access will not display advertising to enrolled users.")

add_heading(doc, "Commercial conditions", 1)
add_bullet(doc, "Seat band. ", "The selected school licence applies to the agreed student population. A school that exceeds its licensed band will move to the next appropriate option at renewal or through an agreed adjustment.")
add_bullet(doc, "Effective cost per learner. ", "The licence fee is fixed for each enrolment band, so the effective per-learner cost depends on the number of active students. It falls as enrolment approaches the upper limit of the band.")
add_bullet(doc, "Licence validity. ", "Access begins on the agreed start date and remains active until the stated term or session expiry date. Renewal is required to continue premium school access after expiry.")
add_bullet(doc, "Payment and invoicing. ", "The final service agreement and invoice will confirm the payment schedule, applicable taxes, licensed population and start and expiry dates.")
add_bullet(doc, "School-funded or learner-funded access. ", "The school may sponsor the licence or agree an approved learner contribution. Any amount charged to families should be transparent and documented in the school's arrangement with Tech Solution Providers Ltd.")
add_bullet(doc, "Items outside the standard fee. ", "Devices, internet or mobile data, extensive on-site training, custom integrations, special data migration and school-specific software development are not included unless expressly stated. Any such work will be scoped and priced before it begins.")

add_heading(doc, "Choosing the right option", 2)
add_text(doc, "The school should select a licence based on the number of students expected to use Quiks during the paid period, rather than total school enrolment where only selected classes will participate. Tech Solution Providers Ltd will confirm the appropriate band during scoping so the proposal, invoice and licence record use the same enrolment basis.")

# Page 7 usage
add_page_break(doc)
add_section_marker(doc, "4", "School use and implementation")
add_heading(doc, "How a school can use Quiks", 1)
add_number(doc, "Define the learning focus. ", "The school identifies the classes, subjects, curriculum or examination priorities and the teachers responsible for the pilot or rollout.")
add_number(doc, "Configure the school account. ", "The administrator selects enrolment fields, chooses shared or individual codes, enrols staff and students, and confirms authorised roles.")
add_number(doc, "Create or connect classrooms. ", "The administrator may create classes centrally and issue class codes, or teachers may create classrooms and submit their codes for linkage to the school portal.")
add_number(doc, "Publish learning materials and activities. ", "Teachers upload lesson notes, set read or download permissions, create tests and assignments, review questions and set submission deadlines.")
add_number(doc, "Students participate. ", "Learners read notes, complete assigned work, practise independently, use explanations and open the Learning Hub when a topic needs more attention.")
add_number(doc, "Review and respond. ", "Teachers review classroom results, while authorised administrators use the central results view for school-level follow-up and reporting workflows.")

add_heading(doc, "Suggested uses across the school", 2)
use_table = doc.add_table(rows=1, cols=3)
use_table.alignment = WD_TABLE_ALIGNMENT.CENTER
use_table.autofit = False
uw = [1.35, 3.15, 2.00]
for i, h in enumerate(("Use case", "How Quiks supports it", "Suggested routine")):
    set_width(use_table.cell(0, i), uw[i])
    set_cell_shading(use_table.cell(0, i), NAVY)
    add_table_text(use_table.cell(0, i), h, bold=True, color=WHITE, size=8.8, align=WD_ALIGN_PARAGRAPH.CENTER)
set_repeat_table_header(use_table.rows[0])
for ri, data in enumerate((
    ("Homework and consolidation", "Topic-focused practice, lesson notes, explanations and teacher-created assignments", "Set one focused activity after a taught unit"),
    ("Assessment and CBT", "Teacher-reviewed questions, deadlines, submissions and objective-question scoring", "Run continuous assessment or CBT practice under school supervision"),
    ("Exam preparation", "Grade, curriculum and target-examination context with repeatable practice", "Schedule weekly revision in priority subjects"),
    ("Intervention", "Additional practice and recorded activity for learners who need more support", "Assign focused follow-up after reviewing results"),
    ("Enrichment", "Learning Hub, direct questions and Competition Arena", "Use in clubs, holiday learning and extension programmes"),
), start=1):
    cells = use_table.add_row().cells
    keep_row_together(use_table.rows[-1])
    for ci, (value, width) in enumerate(zip(data, uw)):
        set_width(cells[ci], width)
        set_cell_shading(cells[ci], WHITE if ri % 2 else PALE_BLUE)
        add_table_text(cells[ci], value, bold=(ci == 0), size=8.5)
set_table_borders(use_table)

# Page 8 pilot
add_page_break(doc)
add_section_marker(doc, "5", "Pilot and partnership recommendation")
add_heading(doc, "Start with a focused school pilot", 1)
add_text(doc, "We recommend a four-week pilot with selected classes and teacher champions. The pilot should test usability, curriculum fit, teacher workflow, learner participation and the practical value of Quiks before a broader rollout. Where necessary, the parties may agree an evaluation period of up to eight weeks.")

pilot = doc.add_table(rows=5, cols=2)
pilot.alignment = WD_TABLE_ALIGNMENT.CENTER
pilot.autofit = False
pw = [1.55, 4.95]
pilot_rows = [
    ("Recommended scope", "One year group or selected classes, two to four subjects and nominated teacher coordinators."),
    ("Onboarding", "School account setup, administrator guidance, enrolment configuration, classroom setup and teacher orientation."),
    ("Success measures", "Account activation, participation, activity completion, teacher usability feedback and learner engagement feedback."),
    ("School commitment", "A decision-maker, teacher champions, access to the selected pilot group and structured feedback at agreed checkpoints."),
    ("Decision after pilot", "The school may end the pilot without a licence commitment or proceed to an agreed term or session package."),
]
for ri, (label, body) in enumerate(pilot_rows):
    keep_row_together(pilot.rows[ri])
    set_width(pilot.cell(ri, 0), pw[0])
    set_width(pilot.cell(ri, 1), pw[1])
    set_cell_shading(pilot.cell(ri, 0), PURPLE)
    set_cell_shading(pilot.cell(ri, 1), WHITE if ri % 2 == 0 else PALE)
    add_table_text(pilot.cell(ri, 0), label, bold=True, color=WHITE, size=8.7)
    add_table_text(pilot.cell(ri, 1), body, size=9.0)
set_table_borders(pilot)

add_heading(doc, "Responsible adoption", 2)
add_bullet(doc, "Teacher-guided use. ", "Quiks should follow the school's scheme of work, approved resources and professional judgement.")
add_bullet(doc, "Account continuity. ", "Students and staff should use their assigned signed-in accounts so authorised school membership, classroom records and licence access remain connected.")
add_bullet(doc, "Review before publication. ", "Teachers should review AI-assisted lesson notes and questions before students use them.")
add_bullet(doc, "Age-appropriate supervision. ", "The school and parents should agree expectations for devices, screen time, acceptable use and communication within class spaces.")
add_bullet(doc, "Data protection arrangements. ", "The parties should complete the applicable service, privacy and data-processing documentation before full deployment.")

add_heading(doc, "Agreements", 2)
add_text(doc, "Before the pilot, the parties may enter a confidentiality agreement where required. If the school proceeds after the pilot, the parties will document the selected licence, service scope, support arrangements, data-processing responsibilities and any school-specific requirements in the appropriate agreement.")

# Page 9 next steps/contact
add_page_break(doc)
add_section_marker(doc, "6", "Next steps")
add_heading(doc, "Proposed next step", 1)
add_text(doc, "We propose a 30-minute demonstration and scoping meeting with the school's decision-maker and nominated teacher representatives. The meeting will confirm the relevant Quiks variant, pilot classes, priority subjects, enrolment method, success measures and the commercial option to be considered after the pilot.")

add_heading(doc, "Information to confirm during scoping", 2)
add_bullet(doc, "Pilot population. ", "The classes, year groups and estimated number of students who will participate.")
add_bullet(doc, "Academic scope. ", "Priority subjects, curriculum requirements and any target examinations.")
add_bullet(doc, "School roles. ", "The decision-maker, school administrator and teacher champions responsible for the rollout.")
add_bullet(doc, "Technology readiness. ", "Available devices, internet access and whether learners will use school or personal devices.")
add_bullet(doc, "Assessment use. ", "Whether the school plans to use Quiks for homework, continuous assessment, CBT practice or another approved purpose.")
add_bullet(doc, "Commercial preference. ", "School-funded access, per-learner access or the appropriate school licence band.")

add_heading(doc, "Contact", 1)
contact = doc.add_table(rows=4, cols=2)
contact.alignment = WD_TABLE_ALIGNMENT.LEFT
contact.autofit = False
cw = [1.55, 4.95]
contact_rows = [
    ("Organisation", "Tech Solution Providers Ltd"),
    ("Product", "Quiks School"),
    ("Representative", "[NAME AND TITLE]"),
    ("Email and telephone", "quiks@techsolutionproviders.net  |  [TELEPHONE NUMBER]"),
]
for ri, (label, value) in enumerate(contact_rows):
    keep_row_together(contact.rows[ri])
    set_width(contact.cell(ri, 0), cw[0])
    set_width(contact.cell(ri, 1), cw[1])
    set_cell_shading(contact.cell(ri, 0), NAVY)
    set_cell_shading(contact.cell(ri, 1), WHITE if ri % 2 == 0 else PALE_BLUE)
    add_table_text(contact.cell(ri, 0), label, bold=True, color=WHITE, size=9.0)
    add_table_text(contact.cell(ri, 1), value, bold=(ri == 0), size=9.5)
set_table_borders(contact)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(18)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Quiks  Learn fast  Grow steadily")
set_font(r, size=10, bold=True, color=PURPLE)

# Metadata
props = doc.core_properties
props.title = "Quiks School Partnership Proposal"
props.subject = "School partnership, implementation and pricing proposal"
props.author = "Tech Solution Providers Ltd"
props.keywords = "Quiks School, education, classroom, CBT, school licence"

doc.save(OUT)
print(OUT)
