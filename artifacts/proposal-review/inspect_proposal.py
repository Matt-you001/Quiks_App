from __future__ import annotations

import hashlib
import json
import sys
from pathlib import Path

from docx import Document


source = Path(sys.argv[1])
document = Document(source)

paragraphs = []
for index, paragraph in enumerate(document.paragraphs):
    text = paragraph.text.strip()
    if text:
        paragraphs.append({"index": index, "style": paragraph.style.name, "text": text})

tables = []
for table_index, table in enumerate(document.tables):
    rows = []
    for row in table.rows:
        rows.append(["\n".join(p.text for p in cell.paragraphs).strip() for cell in row.cells])
    tables.append({"index": table_index, "rows": rows})

headers = []
footers = []
for section_index, section in enumerate(document.sections):
    headers.append({"section": section_index, "paragraphs": [p.text for p in section.header.paragraphs if p.text.strip()]})
    footers.append({"section": section_index, "paragraphs": [p.text for p in section.footer.paragraphs if p.text.strip()]})

digest = hashlib.sha256()
with source.open("rb") as handle:
    for chunk in iter(lambda: handle.read(1024 * 1024), b""):
        digest.update(chunk)

print(json.dumps({
    "source": str(source),
    "sha256": digest.hexdigest(),
    "section_count": len(document.sections),
    "paragraphs": paragraphs,
    "tables": tables,
    "headers": headers,
    "footers": footers,
}, indent=2, ensure_ascii=False))
