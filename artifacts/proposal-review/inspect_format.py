from __future__ import annotations

import json
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


source = Path(sys.argv[1])
doc = Document(source)

paragraph_samples = []
for index, paragraph in enumerate(doc.paragraphs):
    if not paragraph.text.strip():
        continue
    fmt = paragraph.paragraph_format
    runs = []
    for run in paragraph.runs:
        if not run.text:
            continue
        color = run.font.color.rgb
        runs.append({
            "text": run.text[:80],
            "font": run.font.name,
            "size_pt": run.font.size.pt if run.font.size else None,
            "bold": run.bold,
            "italic": run.italic,
            "color": str(color) if color else None,
        })
    paragraph_samples.append({
        "index": index,
        "style": paragraph.style.name,
        "text": paragraph.text[:160],
        "alignment": str(paragraph.alignment),
        "space_before_pt": fmt.space_before.pt if fmt.space_before else None,
        "space_after_pt": fmt.space_after.pt if fmt.space_after else None,
        "line_spacing": fmt.line_spacing,
        "keep_with_next": fmt.keep_with_next,
        "page_break_before": fmt.page_break_before,
        "runs": runs,
    })

tables = []
for index, table in enumerate(doc.tables):
    grid = table._tbl.tblGrid
    widths = [int(col.get(qn("w:w"))) for col in grid.gridCol_lst] if grid is not None else []
    tables.append({
        "index": index,
        "rows": len(table.rows),
        "cols": len(table.columns),
        "style": table.style.name if table.style else None,
        "grid_twips": widths,
    })

with zipfile.ZipFile(source) as archive:
    parts = []
    for info in archive.infolist():
        parts.append({"path": info.filename, "size": info.file_size, "crc": info.CRC})

print(json.dumps({"paragraph_samples": paragraph_samples, "tables": tables, "parts": parts}, indent=2))
