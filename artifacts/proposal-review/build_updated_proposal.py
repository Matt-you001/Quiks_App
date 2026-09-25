from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\USER\Desktop\Quiks")
SOURCE = ROOT / "artifacts" / "proposal-review" / "reference-copy.docx"
OUTPUT = ROOT / "artifacts" / "Quiks_School_Partnership_Updated_Proposal.docx"

PURPLE = "6F2DBD"
NAVY = "1F3D59"
TEXT = "24364B"
MUTED = "5B6B7A"
PALE_BLUE = "EEF4F8"
PALE_PURPLE = "F3EFF9"
WHITE = "FFFFFF"
BORDER = "D8E0E8"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def replace_text_in_runs(container, old, new):
    for paragraph in container.paragraphs:
        for run in paragraph.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)


def remove_body_from(doc, paragraph):
    body = doc.element.body
    start = paragraph._element
    deleting = False
    for child in list(body):
        if child is start:
            deleting = True
        if deleting and child.tag != qn("w:sectPr"):
            body.remove(child)


def format_run(run, size=11, bold=False, color=TEXT, font="Aptos"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def format_paragraph(paragraph, before=0, after=6, line=1.08, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep
    fmt.widow_control = True


def add_kicker(doc, text, new_page=True):
    p = doc.add_paragraph()
    if new_page:
        p.paragraph_format.page_break_before = True
    format_paragraph(p, after=7, keep=True)
    run = p.add_run(text.upper())
    format_run(run, size=9, bold=True, color=PURPLE)
    run.font.all_caps = True
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    format_paragraph(p, before=0 if level == 1 else 5, after=7, keep=True)
    run = p.add_run(text)
    format_run(run, size=20 if level == 1 else 14, bold=True, color="111111", font="Aptos Display")
    return p


def add_body(doc, text, after=7, italic=False):
    p = doc.add_paragraph()
    format_paragraph(p, after=after)
    run = p.add_run(text)
    format_run(run)
    run.font.italic = italic
    return p


def add_bullet(doc, lead, text, numbered=False):
    p = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
    format_paragraph(p, after=4, line=1.04)
    r1 = p.add_run(lead)
    format_run(r1, bold=True)
    r2 = p.add_run(text)
    format_run(r2)
    return p


def add_note(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.7)
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE_PURPLE)
    set_cell_border(cell, color="CFC2E8", size="8")
    set_cell_margins(cell, top=130, start=150, bottom=130, end=150)
    p = cell.paragraphs[0]
    format_paragraph(p, after=0)
    r1 = p.add_run(title + " ")
    format_run(r1, bold=True, color=PURPLE)
    r2 = p.add_run(text)
    format_run(r2)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_table(doc, headers, rows, widths=None, header_color=NAVY, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = 1
    table.style = "Table Grid"
    if widths:
        for i, width in enumerate(widths):
            table.columns[i].width = Inches(width)

    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, header_color)
        set_cell_margins(cell, top=100, start=100, bottom=100, end=100)
        set_cell_border(cell, color=header_color, size="8")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        format_paragraph(p, after=0)
        r = p.add_run(value)
        format_run(r, size=font_size, bold=True, color=WHITE)

    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        fill = PALE_BLUE if row_index % 2 == 0 else WHITE
        for i, value in enumerate(values):
            cell = row.cells[i]
            set_cell_shading(cell, fill)
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            format_paragraph(p, after=0, line=1.0)
            r = p.add_run(value)
            format_run(r, size=font_size, bold=(i == 0), color=TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


doc = Document(SOURCE)

# Preserve the complete first-page cover and its embedded product artwork.
replace_text_in_runs(doc.tables[1].cell(0, 2), "QUIKS UNI", "QUIKS ADVANCE")
doc.tables[0].cell(2, 1).text = "25 SEPTEMBER 2026"
for paragraph in doc.tables[0].cell(2, 1).paragraphs:
    for run in paragraph.runs:
        format_run(run, size=10, bold=True, color=TEXT)

doc.paragraphs[1].text = "An integrated academic, assessment and school-operations partnership for students, teachers and school leadership"
for run in doc.paragraphs[1].runs:
    format_run(run, size=15, color=PURPLE, font="Aptos Display")

doc.paragraphs[5].text = (
    "To strengthen teaching, assessment, reporting and learning beyond lesson time, while giving authorised school "
    "leaders practical tools for academic oversight and selected administrative workflows."
)
for run in doc.paragraphs[5].runs:
    format_run(run, size=11, color=TEXT)

# Replace the previous body while retaining the reference cover, theme, section and footer.
remove_body_from(doc, doc.paragraphs[7])

# 1. Overview
add_kicker(doc, "1  |  Overview", new_page=False)
add_heading(doc, "Why Quiks")
add_body(
    doc,
    "Quiks is a connected learning, assessment and school-support platform for different stages of education. "
    "Quiks Children supports younger learners, Quiks Teens serves secondary-school and college-age students, and "
    "Quiks Advance supports tertiary and advanced learners. A school may adopt the variant or combination that fits "
    "its age groups, curricula and examination priorities.",
)
add_body(
    doc,
    "The platform complements classroom teaching rather than replacing it. Teachers retain responsibility for curriculum "
    "delivery, question approval, marking of written work and interpretation of results. School-licensed access is available "
    "on supported web and mobile devices and is provided without advertising.",
)
add_heading(doc, "What the school receives", level=2)
add_bullet(doc, "Connected school and classroom workflows. ", "Administrators can enrol users, connect classes and maintain central academic records, while teachers manage day-to-day classroom delivery.")
add_bullet(doc, "Flexible curriculum configuration. ", "A school may select one or more curricula for hybrid delivery. Those selections guide school-linked class activities; independent learners continue to use the curriculum in their own profiles.")
add_bullet(doc, "School-appropriate terminology. ", "The administrator can configure the grade, year or class labels offered when teachers set activities, while teachers remain free to name their classrooms as they choose.")
add_bullet(doc, "Assessment and CBT support. ", "Teachers can prepare objective, written or mixed activities, assign marks, set deadlines, include labelled images and control how completed results reach the school portal.")
add_bullet(doc, "Central results and reporting. ", "Authorised staff can review linked classes, student rosters, teacher-submitted result sets and individual student reports.")
add_bullet(doc, "Online and offline delivery options. ", "Schools can run connected activities, prepare signed offline exam packages for Quiks Exam Player, or use a local-network Exam Hub where internet access is unreliable.")
add_heading(doc, "A practical role in school delivery", level=2)
add_body(doc, "Quiks can support homework, revision, intervention, enrichment, lesson-note distribution, continuous assessment, selected CBT exercises, examinations and school reporting. The school chooses the scale, modules and operating rules that fit its policies and technology environment.")

# 2. Academic capabilities
add_kicker(doc, "2  |  Academic and assessment capabilities")
add_heading(doc, "Teaching, classrooms and lesson content")
add_bullet(doc, "Flexible classroom creation. ", "A school administrator may create classes centrally and issue class codes, or teachers may create classes and submit their codes for linkage to the school portal.")
add_bullet(doc, "Lesson-note workflows. ", "Teachers can upload notes, preserve available source illustrations, request different levels of refinement, add necessary labelled illustrations, review the result and publish it as read-only or downloadable content.")
add_bullet(doc, "Document and image support. ", "Teachers can create custom questions with labelled images and may import supported documents for question extraction, subject to the quality and structure of the source material.")
add_bullet(doc, "Class communication. ", "Teachers and enrolled students can communicate in the shared classroom space under the school’s acceptable-use rules.")
add_heading(doc, "Assessment and CBT", level=2)
add_bullet(doc, "Objective, written and mixed activities. ", "Objective items are scored automatically. Teachers mark written responses, and students are informed when the final score depends on teacher review.")
add_bullet(doc, "Question-level control. ", "Teachers can set marks for individual questions, review generated content and marking guidance, edit questions and publish only when satisfied.")
add_bullet(doc, "Candidate navigation. ", "Students may skip a question and return after the unskipped questions, reducing avoidable time loss during an activity.")
add_bullet(doc, "Configurable exit policy. ", "Teachers can choose warn-and-record, auto-submit on confirmed exit, or strict CBT mode that records prohibited background or tab changes and submits the attempt.")
add_bullet(doc, "Broad subject coverage. ", "A General question focus is designed to sample across the applicable topic range for the selected subject rather than repeatedly narrowing activity generation to one topic.")
add_note(doc, "Teacher review remains essential.", "AI-assisted questions, answers, diagrams and lesson-note refinements should be checked for curriculum fit, factual accuracy, accessibility and age appropriateness before release.")

# 3. Results and reporting
add_kicker(doc, "3  |  Results, reports and school visibility")
add_heading(doc, "Teacher-controlled publication")
add_body(doc, "Student submissions first remain within the classroom workflow. The class teacher decides which completed activity results are formally published to the school portal, preventing informal practice or draft activities from automatically becoming part of the school’s central record.")
add_heading(doc, "Central school results register", level=2)
add_bullet(doc, "Class-level visibility. ", "An authorised administrator can open a linked classroom, view the enrolled student list and open the results submitted for that class.")
add_bullet(doc, "Activity-title columns. ", "Result tables use the title supplied by the teacher for each published activity rather than imposing generic Test, Assignment or Examination headings.")
add_bullet(doc, "Simplified review. ", "Schools can filter results by subject and date, then print or export the resulting table for approved internal use.")
add_bullet(doc, "Individual report sheets. ", "The reporting workflow can prepare a student-specific report with spaces for the class teacher’s and Principal or Head Teacher’s signatures.")
add_bullet(doc, "Guardian communication. ", "Approved reports can be printed, exported and sent to a verified student or guardian email address, subject to the school’s reporting policy and final review.")
add_heading(doc, "Learning Hub and independent study", level=2)
add_body(doc, "The Learning Hub separates three clear workflows: Generate Lesson, Past Q&A and Ask a Question. Past Q&A can accept pasted text or supported question documents, extract questions, provide worked answers and search consented library material by examination, subject and year.")
add_body(doc, "Where the contributor confirms permission to share, extracted questions and answers may support the Quiks past-question library and improve exam-aligned retrieval and generation context. Quiks does not claim that each upload directly retrains the underlying foundation model.")

# 4. Offline examinations
add_kicker(doc, "4  |  Offline examination options")
add_heading(doc, "Assessment where internet access is limited")
add_body(doc, "Teachers prepare and approve an activity online, then export a protected exam package. Student packages exclude answer keys. A separate encrypted teacher marking guide is provided only to authorised staff.")
add_table(
    doc,
    ["Delivery option", "How it works", "Result handling"],
    [
        ["Quiks Exam Player", "The player is built into supported Quiks apps and opens signed offline exam packages on student devices.", "Attempts are saved locally and may be synchronised when connectivity returns."],
        ["Standalone offline mode", "The school distributes a protected package for use on approved devices without requiring later connection to Quiks.", "Responses are exported locally; the school marks and computes results under its own process."],
        ["Quiks Exam Hub", "A lightweight local service distributes activities across a school laboratory or local network without public internet.", "The hub collects local attempts and produces approved exports for the school."],
        ["Printable paper mode", "The school exports a print-ready paper and a separately controlled marking guide.", "Marking and result computation remain fully with the school."],
    ],
    widths=[1.45, 3.15, 2.1],
    header_color=PURPLE,
    font_size=9,
)
add_heading(doc, "Operational safeguards", level=2)
add_bullet(doc, "Package integrity. ", "Offline packages are signed so the player can detect unauthorised modification before an activity starts.")
add_bullet(doc, "Candidate privacy. ", "Student packages do not contain answer keys or the teacher marking guide.")
add_bullet(doc, "Local resilience. ", "Responses are autosaved during supported offline sessions to reduce loss from power or device interruptions.")
add_bullet(doc, "Appropriate deployment. ", "Native app or local-network deployment is recommended for supervised examinations. Ordinary browser/PWA delivery offers weaker lockdown and should be treated as a secondary option.")

# 5. Administration
add_kicker(doc, "5  |  School administration capabilities")
add_heading(doc, "Optional operational support for schools")
add_body(doc, "Alongside the academic platform, Quiks includes school-operations modules that can be enabled for an institution according to its needs, readiness and authorised roles. Schools that do not need these modules can continue using the academic features with the essential administration already required for enrolment, classrooms and results.")
add_table(
    doc,
    ["Module", "Current capabilities"],
    [
        ["Operations foundation", "Student, staff and guardian registry; names, class, admission number, email, registration date, configurable fields, archive controls and authorised exports."],
        ["Attendance", "Daily attendance capture, approved amendments with reasons, and traceable changes for authorised school leadership."],
        ["Planning and scheduling", "Lesson planning plus lesson and examination timetables for selected school workflows."],
        ["Staff management", "Restricted staff reports and records accessible according to assigned roles."],
        ["Transport management", "Route, vehicle and passenger records; uniform pricing across routes or required route-by-route prices when varying pricing is selected."],
        ["Governance", "Role-based access, school-scoped records, audit history and data export for authorised administrators."],
    ],
    widths=[1.65, 5.05],
    header_color=NAVY,
    font_size=9.2,
)
add_heading(doc, "School choice and permissions", level=2)
add_body(doc, "The school decides whether student or staff photographs, birth certificates or identity documents, and medical documents are collected. These categories remain disabled unless the school administrator deliberately enables them, and collection should be limited to a documented purpose.")
add_bullet(doc, "Create other administrators: ", "School Owner.")
add_bullet(doc, "Assign or remove roles; edit submitted attendance; view disciplinary or appraisal records: ", "School Owner or Administrator.")
add_bullet(doc, "Export all school data: ", "Administrator.")
add_bullet(doc, "Delete or archive student and staff records: ", "School Owner, with confirmation and audit history.")

# 6. Value to stakeholders - school and teachers
add_kicker(doc, "6  |  Value to stakeholders")
add_heading(doc, "Value to the school")
add_bullet(doc, "One connected academic view. ", "School enrolment, curricula, class links, teacher-published results and approved reports can be reviewed from a central portal.")
add_bullet(doc, "A practical CBT pathway. ", "The school can combine online assessment, supervised app delivery, local-network delivery and printable fallbacks according to available infrastructure.")
add_bullet(doc, "Consistent school standards. ", "Administrators can configure curricula, activity-level terminology, enrolment fields and selected administration modules for school-linked use.")
add_bullet(doc, "Improved reporting evidence. ", "Filtered class results, individual report sheets, exports and signature workflows support follow-up and parent communication.")
add_bullet(doc, "Operational flexibility. ", "Schools can begin with academic delivery and enable only the administration modules they require during the introductory period.")
add_bullet(doc, "Scalable access. ", "Licence bands support selected learners, small schools, growing schools, larger enrolments and multi-campus networks.")
add_heading(doc, "Value to teachers")
add_bullet(doc, "Faster preparation with control. ", "Teachers can generate or create activities, import supported content, allocate marks and review every question before publication.")
add_bullet(doc, "More suitable assessment formats. ", "Objective, written and mixed activities support routine exercises, assignments, tests and examinations.")
add_bullet(doc, "Reduced repetitive marking. ", "Objective responses are scored automatically while teachers retain professional responsibility for written responses and final review.")
add_bullet(doc, "Selective school reporting. ", "Teachers decide which completed activity results become part of the school’s central results register.")
add_bullet(doc, "Reusable teaching resources. ", "Lesson notes, diagrams, past-question workflows, class communication and curriculum-aware support reduce duplication across the teaching cycle.")

# 7. Value to students and parents
add_kicker(doc, "6  |  Value to stakeholders continued")
add_heading(doc, "Value to students")
add_bullet(doc, "Guided independent study. ", "Students can practise broadly across a subject or focus on selected topics at an appropriate level and curriculum context.")
add_bullet(doc, "Clear feedback. ", "Automatic scoring, explanations and teacher-marked written responses help learners identify what they understand and where further work is needed.")
add_bullet(doc, "Exam-focused preparation. ", "Past Q&A, teacher-set activities and offline exam delivery support preparation for local and international examinations.")
add_bullet(doc, "Learning continuity. ", "Lesson notes, assignments, tests, class communication and selected offline packages extend learning beyond a single lesson or internet connection.")
add_bullet(doc, "Study habit tracking. ", "Daily targets include time spent in eligible practice, tests and assignments, giving the learner a more complete view of study activity.")
add_bullet(doc, "Motivation and recognition. ", "Competition and Top Performers views can recognise achievement and display a learner’s school where a school name is recorded.")
add_heading(doc, "Value to parents and guardians")
add_bullet(doc, "A structured home-study route. ", "Quiks organises practice, revision, published lesson notes and assigned work outside school hours.")
add_bullet(doc, "Better school-home continuity. ", "Approved student reports can be shared with verified guardians after teacher or school review.")
add_bullet(doc, "Clearer conversations. ", "Study time, activity completion and reviewed performance records can support practical discussions about effort, strengths and intervention needs.")
add_bullet(doc, "Cost-effective school access. ", "Per-learner and school packages may offer a lower-cost route than separate individual subscriptions, depending on the selected package and participation level.")
add_note(doc, "Shared responsibility.", "Quiks does not replace teacher judgement, school safeguarding or parental supervision. Schools and families should agree appropriate expectations for devices, screen time, acceptable use and AI-assisted learning.")

# 8. Pricing
add_kicker(doc, "7  |  Financial implications")
add_heading(doc, "Academic licence options")
add_body(doc, "The school may begin with a four-week pilot at no licence cost. An extension of up to eight weeks may be agreed where the school calendar or evaluation plan requires more time. If the school does not proceed after the agreed pilot, no onboarding or licence fee is payable.")
add_body(doc, "If the school adopts Quiks, a one-time onboarding fee of NGN 50,000 becomes payable. The school then selects the academic licence that matches its enrolled population and preferred payment cycle.")
add_table(
    doc,
    ["Licence option", "Eligible users", "Monthly fee", "Per term", "Per session"],
    [
        ["Per Learner Access", "Individual students enrolled through the school", "NGN 1,500 per learner", "Approx. NGN 4,500 per learner", "NGN 12,500"],
        ["Essential School", "10 to 100 students", "Not applicable", "NGN 300,000", "NGN 750,000"],
        ["Growth School", "101 to 200 students", "Not applicable", "NGN 500,000", "NGN 1,250,000"],
        ["Comprehensive School", "201 to 500 students", "Not applicable", "NGN 800,000", "NGN 1,950,000"],
        ["Enterprise Network", "School group or multiple campuses; scope confirmed in the agreement", "Not applicable", "NGN 1,000,000", "NGN 2,500,000"],
    ],
    widths=[1.28, 2.08, 1.15, 1.13, 1.13],
    header_color=PURPLE,
    font_size=8.5,
)
add_heading(doc, "Session payment savings", level=2)
add_body(doc, "A session licence costs less than purchasing three separate term licences:")
add_table(
    doc,
    ["Licence option", "Three terms", "Session fee", "Saving"],
    [
        ["Essential School", "NGN 900,000", "NGN 750,000", "NGN 150,000"],
        ["Growth School", "NGN 1,500,000", "NGN 1,250,000", "NGN 250,000"],
        ["Comprehensive School", "NGN 2,400,000", "NGN 1,950,000", "NGN 450,000"],
        ["Enterprise Network", "NGN 3,000,000", "NGN 2,500,000", "NGN 500,000"],
    ],
    widths=[2.25, 1.48, 1.48, 1.48],
    header_color=NAVY,
    font_size=9.2,
)

# 9. Commercial scope
add_kicker(doc, "7  |  Financial implications continued")
add_heading(doc, "What the academic licence covers")
add_bullet(doc, "Platform access. ", "Use of the relevant Quiks variant on supported web and mobile devices for the licensed users and validity period.")
add_bullet(doc, "Core school controls. ", "School enrolment, authorised roles, class linkage, curricula, licence information and the central academic results workflow.")
add_bullet(doc, "Teaching and assessment. ", "Classrooms, lesson notes, tests, assignments, class chat, practice, Learning Hub, Competition Arena and supported offline assessment exports.")
add_bullet(doc, "Initial onboarding. ", "Account setup guidance, administrator orientation, enrolment configuration and a short teacher introduction, funded through the one-time onboarding fee.")
add_bullet(doc, "Advertising-free school use. ", "School-licensed access does not display advertising to enrolled users.")
add_heading(doc, "Complimentary administration modules during the introductory period", level=2)
add_note(doc, "Current commercial position.", "A school with an active paid academic licence may request any currently available Quiks school-administration module to be enabled at no additional licence charge during the introductory period. Activation is subject to agreed scope, setup, readiness and product availability.")
add_body(doc, "No separate Basic or Advance administration package is being priced in this proposal. Tech Solution Providers Ltd may introduce grouped administration packages in the future. Any future pricing or packaging change will be communicated in advance and agreed for a subsequent renewal; it will not alter the school’s current contracted licence period unless both parties agree in writing.")
add_heading(doc, "Commercial conditions", level=2)
add_bullet(doc, "Licence validity. ", "Access begins on the agreed start date and remains active until the stated term or session expiry date. Renewal is required to continue premium school access after expiry.")
add_bullet(doc, "Seat band. ", "A school that exceeds its licensed student band moves to the next appropriate option at renewal or through an agreed adjustment.")
add_bullet(doc, "Payment and invoicing. ", "The final agreement and invoice confirm payment, applicable taxes, licensed population, validity dates and any approved additional work.")
add_bullet(doc, "Items outside the standard fee. ", "Devices, connectivity, extensive on-site training, custom integrations, special migration and school-specific development are separately scoped and priced before work begins.")

# 10. Implementation
add_kicker(doc, "8  |  School use and implementation")
add_heading(doc, "A controlled rollout")
add_bullet(doc, "Define the scope. ", "Confirm participating classes, subjects, curricula, examinations, administration modules and success measures.", numbered=True)
add_bullet(doc, "Configure the school. ", "Set enrolment fields, collection toggles, roles, grade terminology, curriculum selections and shared or individual enrolment codes.", numbered=True)
add_bullet(doc, "Connect classrooms. ", "Create classes centrally or link teacher-created classrooms to the school portal.", numbered=True)
add_bullet(doc, "Prepare staff. ", "Orient administrators and teacher champions on lesson, activity, marking, result-publication and reporting workflows.", numbered=True)
add_bullet(doc, "Enrol students and deliver. ", "Students join with approved credentials or codes and use the relevant online, app or offline pathway.", numbered=True)
add_bullet(doc, "Review and improve. ", "Monitor participation, teacher experience, result quality, safeguarding, technical performance and support needs.", numbered=True)
add_heading(doc, "Suggested uses across the school", level=2)
add_table(
    doc,
    ["Use case", "How Quiks supports it", "Suggested routine"],
    [
        ["Homework and consolidation", "Topic-focused practice, lesson notes, explanations and teacher-created assignments", "One focused activity after a taught unit"],
        ["Assessment and CBT", "Objective, written or mixed activities; marks, deadlines, controls and reviewed results", "Continuous assessment or supervised CBT"],
        ["Reporting", "Teacher-selected publication, filtered result tables and individual reports", "Review and release on the school calendar"],
        ["Exam preparation", "Past Q&A, target-exam context and repeatable online or offline practice", "Weekly revision in priority subjects"],
        ["School operations", "Selected registry, attendance, planning, staff or transport modules", "Enable only approved modules and roles"],
    ],
    widths=[1.45, 3.2, 2.05],
    header_color=NAVY,
    font_size=8.8,
)

# 11. Pilot
add_kicker(doc, "9  |  Pilot and partnership recommendation")
add_heading(doc, "Start with a focused school pilot")
add_body(doc, "We recommend a four-week pilot with selected classes and teacher champions. The pilot should test usability, curriculum fit, assessment workflow, teacher-controlled reporting, learner participation and the practical value of any requested administration modules before a broader rollout. Where necessary, the parties may agree an evaluation period of up to eight weeks.")
add_table(
    doc,
    ["Pilot area", "Recommended approach"],
    [
        ["Scope", "One year group or selected classes, two to four subjects, nominated teacher champions and only the administration modules needed for evaluation."],
        ["Onboarding", "School account setup, roles, enrolment configuration, classroom linkage, curriculum settings and staff orientation."],
        ["Success measures", "Activation, participation, activity completion, result-publication workflow, usability feedback and learner engagement."],
        ["School commitment", "A decision-maker, authorised administrators, teacher champions, access to the pilot group and structured feedback."],
        ["Decision", "End the pilot without a licence commitment or proceed to an agreed academic term or session package."],
    ],
    widths=[1.55, 5.15],
    header_color=PURPLE,
    font_size=9.2,
)
add_heading(doc, "Responsible adoption", level=2)
add_bullet(doc, "Review before release. ", "Teachers should review AI-assisted notes, questions, answers, diagrams and marking guidance before students use them.")
add_bullet(doc, "Assessment integrity. ", "The school defines supervision, device, identity, permitted-resource and incident-handling rules for each assessment mode.")
add_bullet(doc, "Data minimisation. ", "Collect only personal information and documents required for a documented school purpose, with suitable access and retention controls.")
add_bullet(doc, "Agreement and privacy. ", "Before full deployment, the parties should complete the applicable service, privacy and data-processing documentation, including the responsibilities of each party.")

# 12. Next steps
add_kicker(doc, "10  |  Next steps")
add_heading(doc, "Proposed next step")
add_body(doc, "We propose a 30-minute demonstration and scoping meeting with the school’s decision-maker, authorised administrator and nominated teacher representatives. The meeting will confirm the relevant Quiks variants, pilot population, academic scope, requested administration modules, technology environment, success measures and commercial option.")
add_heading(doc, "Information to confirm during scoping", level=2)
add_bullet(doc, "Pilot population. ", "Participating classes or year groups and estimated student numbers.")
add_bullet(doc, "Academic scope. ", "Priority subjects, hybrid curriculum requirements, target examinations and reporting expectations.")
add_bullet(doc, "School roles. ", "School Owner, administrators, teacher champions and authorised reporting personnel.")
add_bullet(doc, "Technology readiness. ", "Available devices, internet reliability, local-network facilities and the preferred offline-exam option.")
add_bullet(doc, "Administration scope. ", "Registry, attendance, planning, staff reporting, transport or other currently available modules requested for introductory access.")
add_bullet(doc, "Commercial preference. ", "Per-learner access or the appropriate academic school licence band.")
add_heading(doc, "Contact")
add_table(
    doc,
    ["Contact item", "Details"],
    [
        ["Organisation", "Tech Solution Providers Ltd"],
        ["Product", "Quiks School"],
        ["Representative", "ENGR MATTHEW ONAH / CEO, TECH SOLUTION PROVIDERS LTD."],
        ["Email and telephone", "quiks@techsolutionproviders.net  |  09043980140"],
    ],
    widths=[1.55, 5.15],
    header_color=NAVY,
    font_size=9.5,
)
closing = doc.add_paragraph()
closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
format_paragraph(closing, before=10, after=0)
run = closing.add_run("Quiks  |  Learn fast. Grow steadily.")
format_run(run, size=11, bold=True, color=PURPLE)

# Ensure the existing footer remains consistent and its PAGE field updates in Word.
settings = doc.settings._element
update_fields = settings.find(qn("w:updateFields"))
if update_fields is None:
    update_fields = OxmlElement("w:updateFields")
    settings.append(update_fields)
update_fields.set(qn("w:val"), "true")

doc.core_properties.title = "Quiks School Partnership Proposal"
doc.core_properties.subject = "Academic, assessment and school administration partnership proposal"
doc.core_properties.author = "Tech Solution Providers Ltd"
doc.core_properties.comments = "Updated 25 September 2026 to reflect current Quiks School capabilities and introductory administration-module access."

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
