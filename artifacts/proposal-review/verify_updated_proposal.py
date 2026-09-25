import hashlib
import zipfile
from pathlib import Path

from docx import Document

root = Path(r"C:\Users\USER\Desktop\Quiks")
source = root / "artifacts" / "proposal-review" / "reference-copy.docx"
output = root / "artifacts" / "Quiks_School_Partnership_Updated_Proposal.docx"

expected_source_hash = "4380a5ec74edbd40dbf3fb9672197ddf5bce7c0476166d2fca8b0034130a1f80"
actual_source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
assert actual_source_hash == expected_source_hash, (actual_source_hash, expected_source_hash)

with zipfile.ZipFile(output) as package:
    assert package.testzip() is None
    package_part_count = len(package.namelist())

doc = Document(output)
all_text = "\n".join(p.text for p in doc.paragraphs)
all_text += "\n" + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)

assert "QUIKS ADVANCE" in all_text
assert "Quiks Uni" not in all_text
assert "Level 1" not in all_text
assert "Level 2" not in all_text
assert "Level 3" not in all_text
assert "Complimentary administration modules during the introductory period" in all_text
assert "NGN 300,000" in all_text
assert "NGN 2,500,000" in all_text
assert "Quiks Exam Player" in all_text
assert "Teacher-controlled publication" in all_text

print({
    "source_hash_preserved": True,
    "package_parts": package_part_count,
    "paragraphs": len(doc.paragraphs),
    "tables": len(doc.tables),
    "sections": len(doc.sections),
    "output_bytes": output.stat().st_size,
})
